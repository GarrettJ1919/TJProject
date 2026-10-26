script_content = '''"""
Fills the Before/After stat columns in the TJ surgery spreadsheet using the
Baseball-Reference season CSVs in ./br_csv (one file per year, e.g. 2024.csv).

ONLY the pitchers already in the spreadsheet are filled. Nothing is added,
removed, or reordered, and your original file is never modified: results go
to a new *_FILLED.xlsx with your formatting preserved.

Setup (once):   pip install pandas openpyxl pybaseball
Run:            python pull_pitcher_stats.py
"""
import glob
import os
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook

# ---------------- CONFIG ----------------
INPUT_FILE = "TJ_Surgery_2018_2023_Pitchers_Trimmed.xlsx"
INPUT_SHEET = "MLB Pitchers 2018-2023 TJ"
OUTPUT_FILE = "TJ_Surgery_2018_2023_Pitchers_FILLED.xlsx"
CSV_DIR = "br_csv"

YEARS_BEFORE = 3        # seasons before the surgery year (surgery year excluded)
YEARS_AFTER = 3         # seasons after the surgery year
IP_AS_DECIMAL = False   # False -> baseball notation (71.1 = 71 1/3). True -> 71.333
FILL_FGID = True        # try to fill the blank FgID column (needs pybaseball + internet)

NAME_COL = "Player"
SURGERY_DATE_COL = "Surgery Date"
SURGERY_YEAR_COL = "Surgery Year"
FGID_COL = "FgID"

# stat -> (Before column, After column, Excel number format)
OUT = {
    "G":   ("Games - Before", "Games - After", "0"),
    "GS":  ("Games Started - Before", "Games Started - After", "0"),
    "IP":  ("Innings Pitched - Before", "Innings Pitched - After", "0.0"),
    "KBB": ("K-BB% - Before", "K-BB% - After", "0.0%"),
    "ERA": ("ERA - Before", "ERA - After", "0.00"),
    "FIP": ("FIP - Before", "FIP - After", "0.00"),
}
REQUIRED_BR_COLS = ["G", "GS", "IP", "ER", "SO", "BB", "BF", "HR"]
# ----------------------------------------


def norm(s):
    """'José Berríos Jr.' -> 'jose berrios jr' so accents/punctuation don't break matching."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\\s+", " ", "".join(c for c in s.lower() if c.isalnum() or c == " ")).strip()


def ip_to_decimal(ip):
    """BR writes 71 1/3 innings as 71.1 -- convert to a true decimal (71.333)."""
    if pd.isna(ip):
        return 0.0
    ip = float(ip)
    whole = int(ip)
    return whole + round((ip - whole) * 10) / 3


def decimal_to_ip(x):
    outs = round(x * 3)
    return outs // 3 + (outs % 3) / 10


# ---------------- LOAD BR CSVs ----------------
tables = {}
for path in sorted(glob.glob(os.path.join(CSV_DIR, "*.csv"))):
    m = re.search(r"(20\\d\\d)", os.path.basename(path))
    if not m:
        continue
    year = int(m.group(1))
    t = pd.read_csv(path, encoding="utf-8-sig")

    missing_cols = [c for c in REQUIRED_BR_COLS + ["Player"] if c not in t.columns]
    if missing_cols:
        raise SystemExit(f"{path} is missing columns {missing_cols}. Columns found: {list(t.columns)}")

    # drop repeated header rows, blank rows, league-average rows
    t = t[t["Player"].notna() & (t["Player"] != "Player")]
    t = t[~t["Player"].astype(str).str.contains("League Average", case=False)].copy()

    team_col = "Team" if "Team" in t.columns else ("Tm" if "Tm" in t.columns else None)
    t["Player"] = t["Player"].astype(str).str.replace(r"[*#]", "", regex=True).str.strip()
    t["name_norm"] = t["Player"].apply(norm)
    # unique key: BR id when the export has it, otherwise the normalized name
    if "Player-additional" in t.columns:
        t["key"] = t["Player-additional"].fillna(t["name_norm"])
    else:
        t["key"] = t["name_norm"]

    for c in REQUIRED_BR_COLS + ["HBP"]:
        if c in t.columns and c != "IP":
            t[c] = pd.to_numeric(t[c], errors="coerce")
    if "HBP" not in t.columns:
        t["HBP"] = 0.0
    t["HBP"] = t["HBP"].fillna(0)
    t["IP_dec"] = t["IP"].apply(ip_to_decimal)

    # Traded players appear once per team plus a combined "2TM"/"3TM" row.
    # Keep only the combined row so nothing is double-counted.
    if team_col:
        t["is_tm"] = t[team_col].astype(str).str.match(r"^\\dTM$")
    else:
        t["is_tm"] = False
    t = t.sort_values("is_tm", ascending=False).drop_duplicates("key").copy()

    # FIP = (13*HR + 3*(BB+HBP) - 2*SO)/IP + constant; constant is set each season
    # so that league FIP equals league ERA.
    raw_num = 13 * t["HR"] + 3 * (t["BB"] + t["HBP"]) - 2 * t["SO"]
    lg = t["IP_dec"] > 0
    lg_ip = t.loc[lg, "IP_dec"].sum()
    const = t.loc[lg, "ER"].sum() * 9 / lg_ip - raw_num[lg].sum() / lg_ip
    t["FIP_use"] = raw_num / t["IP_dec"].where(t["IP_dec"] > 0) + const
    if "FIP" in t.columns:  # use BR's own FIP if your export happens to have it
        t["FIP_use"] = pd.to_numeric(t["FIP"], errors="coerce").fillna(t["FIP_use"])

    tables[year] = t.set_index("key", drop=False)
    print(f"{year}: {len(t)} pitchers")

if not tables:
    raise SystemExit(f"No CSVs with a 20XX year in the filename were found in '{CSV_DIR}'.")

# name -> possible BR keys, across all years
all_rows = pd.concat(tables.values(), ignore_index=True)
name_to_keys = all_rows.groupby("name_norm")["key"].unique()

# ---------------- OPEN YOUR SPREADSHEET (formatting preserved) ----------------
wb = load_workbook(INPUT_FILE)
ws = wb[INPUT_SHEET]
headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}

needed = [NAME_COL, SURGERY_DATE_COL] + [c for b, a, _ in OUT.values() for c in (b, a)]
absent = [c for c in needed if c not in headers]
if absent:
    raise SystemExit(f"These columns were not found in row 1 of '{INPUT_SHEET}': {absent}")

unmatched = []
resolved = {}  # excel row -> BR key

for r in range(2, ws.max_row + 1):
    name = ws.cell(r, headers[NAME_COL]).value
    if not name:
        continue
    sd = ws.cell(r, headers[SURGERY_DATE_COL]).value
    sy = pd.to_datetime(sd).year if sd else ws.cell(r, headers.get(SURGERY_YEAR_COL, 0) or 1).value
    if not sy or pd.isna(sy):
        unmatched.append({"Row": r, "Player": name, "Period": "-", "Issue": "No surgery date/year"})
        continue
    sy = int(sy)
    windows = {"Before": range(sy - YEARS_BEFORE, sy), "After": range(sy + 1, sy + YEARS_AFTER + 1)}

    keys = list(name_to_keys.get(norm(name), []))
    if len(keys) > 1:  # same name, different players: keep ones who pitched in this player's windows
        win_years = [y for w in windows.values() for y in w]
        keys = [k for k in keys if any(y in tables and k in tables[y].index for y in win_years)]
    if len(keys) != 1:
        issue = "No name match in any CSV" if not keys else f"Ambiguous name, candidates: {keys}"
        unmatched.append({"Row": r, "Player": name, "Period": "ID", "Issue": issue})
        continue
    key = keys[0]
    resolved[r] = key

    for label, years in windows.items():
        rows = [tables[y].loc[key] for y in years if y in tables and key in tables[y].index]
        if not rows:
            unmatched.append({"Row": r, "Player": name, "Period": label,
                              "Issue": f"No seasons found for {list(years)}"})
            if label == "After":
                for stat, (bcol, acol, fmt) in OUT.items():
                    cell = ws.cell(r, headers[acol])
                    cell.value = "NULL"
                    cell.number_format = "@"
            continue
        s = pd.DataFrame(rows)
        ip = s["IP_dec"].sum()
        bf = s["BF"].sum()
        f = s.dropna(subset=["FIP_use"])
        f = f[f["IP_dec"] > 0]
        vals = {
            "G": s["G"].sum(),
            "GS": s["GS"].sum(),
            "IP": ip if IP_AS_DECIMAL else decimal_to_ip(ip),
            "KBB": (s["SO"].sum() - s["BB"].sum()) / bf if bf > 0 else None,
            "ERA": s["ER"].sum() * 9 / ip if ip > 0 else None,
            "FIP": (f["FIP_use"] * f["IP_dec"]).sum() / f["IP_dec"].sum() if len(f) else None,
        }
        for stat, (bcol, acol, fmt) in OUT.items():
            cell = ws.cell(r, headers[bcol if label == "Before" else acol])
            v = vals[stat]
            cell.value = None if v is None or pd.isna(v) else round(float(v), 4)
            cell.number_format = "0.000" if (stat == "IP" and IP_AS_DECIMAL) else fmt

# ---------------- OPTIONAL: fill FgID ----------------
if FILL_FGID and FGID_COL in headers and resolved:
    try:
        from pybaseball import playerid_reverse_lookup
        bb_ids = sorted({k for k in resolved.values() if re.fullmatch(r"[a-z'.]+\\d{2}", str(k))})
        lk = playerid_reverse_lookup(bb_ids, key_type="bbref")
        bb_to_fg = {r_.key_bbref: int(r_.key_fangraphs) for r_ in lk.itertuples()
                    if pd.notna(r_.key_fangraphs) and r_.key_fangraphs not in (-1, "")}
        for r, key in resolved.items():
            if key in bb_to_fg and ws.cell(r, headers[FGID_COL]).value in (None, ""):
                ws.cell(r, headers[FGID_COL]).value = bb_to_fg[key]
    except Exception as e:
        print(f"Skipped FgID fill ({e}). Stats are still filled.")

# ---------------- SAVE ----------------
if "Unmatched - Check Manually" in wb.sheetnames:
    del wb["Unmatched - Check Manually"]
um = wb.create_sheet("Unmatched - Check Manually")
um.append(["Row", "Player", "Period", "Issue"])
for u in unmatched:
    um.append([u["Row"], u["Player"], u["Period"], u["Issue"]])
for col, w in zip("ABCD", (8, 28, 10, 70)):
    um.column_dimensions[col].width = w

wb.save(OUTPUT_FILE)
print(f"\\nDone. {len(resolved)} players matched, {len(unmatched)} issues logged on the "
      f"'Unmatched - Check Manually' tab.\\nSaved to {OUTPUT_FILE}")
'''

with open('pull_pitcher_stats.py', 'w', encoding='utf-8') as f:
    f.write(script_content)

print("Generated pull_pitcher_stats.py")